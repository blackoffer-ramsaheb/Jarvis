# Document RAG functionality with Multi-Format Support
import os
import shutil
import csv
from pathlib import Path
from typing import List, Dict, Any, Optional

from pypdf import PdfReader
import docx
from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_text_splitters import RecursiveCharacterTextSplitter

BASE_DIR = Path(__file__).resolve().parent.parent
DATA_DIR = BASE_DIR / "data"
DOCS_DIR = DATA_DIR / "documents"
VECTOR_DB_PATH = str(DATA_DIR / "vector_db")

_embeddings = None

def get_embeddings():
    """Lazy-load and cache HuggingFace embeddings model."""
    global _embeddings
    if _embeddings is None:
        _embeddings = HuggingFaceEmbeddings(
            model_name="sentence-transformers/all-MiniLM-L6-v2"
        )
    return _embeddings


def extract_pdf_documents(file_path: str) -> List[Document]:
    """Load text documents from a PDF file using PyPDFLoader or PdfReader fallback."""
    try:
        loader = PyPDFLoader(file_path)
        docs = loader.load()
        if any(d.page_content.strip() for d in docs):
            return docs
    except Exception as e:
        print(f"PyPDFLoader error, falling back to PdfReader: {e}")

    reader = PdfReader(file_path)
    docs = []
    filename = Path(file_path).name
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(Document(
                page_content=text,
                metadata={"source": filename, "page": i + 1, "type": "pdf"}
            ))
    return docs


def extract_docx_documents(file_path: str) -> List[Document]:
    """Extract text from Microsoft Word (.docx) files."""
    doc = docx.Document(file_path)
    filename = Path(file_path).name
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    full_text = "\n\n".join(paragraphs)
    if not full_text.strip():
        return []
    return [Document(
        page_content=full_text,
        metadata={"source": filename, "type": "docx"}
    )]


def extract_text_documents(file_path: str) -> List[Document]:
    """Extract text from plain text or markdown files (.txt, .md)."""
    filename = Path(file_path).name
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()
    except UnicodeDecodeError:
        with open(file_path, "r", encoding="latin-1") as f:
            content = f.read()

    if not content.strip():
        return []

    return [Document(
        page_content=content,
        metadata={"source": filename, "type": Path(file_path).suffix.replace(".", "")}
    )]


def extract_csv_documents(file_path: str) -> List[Document]:
    """Extract structured rows from CSV files."""
    filename = Path(file_path).name
    docs = []
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            reader = csv.reader(f)
            header = next(reader, None)
            if not header:
                return []
            header_str = ", ".join(header)
            row_chunks = []
            for row_idx, row in enumerate(reader, start=1):
                row_text = f"Row {row_idx}: " + ", ".join([f"{col}: {val}" for col, val in zip(header, row)])
                row_chunks.append(row_text)
                if len(row_chunks) >= 15:
                    docs.append(Document(
                        page_content=f"Columns: {header_str}\n" + "\n".join(row_chunks),
                        metadata={"source": filename, "type": "csv"}
                    ))
                    row_chunks = []
            if row_chunks:
                docs.append(Document(
                    page_content=f"Columns: {header_str}\n" + "\n".join(row_chunks),
                    metadata={"source": filename, "type": "csv"}
                ))
    except Exception as e:
        print(f"Error parsing CSV file {file_path}: {e}")
    return docs


def load_file_documents(file_path: str) -> List[Document]:
    """Dispatch file extraction based on extension."""
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return extract_pdf_documents(file_path)
    elif ext in [".docx", ".doc"]:
        return extract_docx_documents(file_path)
    elif ext in [".txt", ".md", ".log", ".json"]:
        return extract_text_documents(file_path)
    elif ext == ".csv":
        return extract_csv_documents(file_path)
    else:
        # Default text attempt
        return extract_text_documents(file_path)


def rebuild_vector_db_from_scratch() -> int:
    """Scan all files in data/documents and rebuild FAISS index."""
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    all_chunks = []
    splitter = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=120)

    for doc_file in DOCS_DIR.iterdir():
        if doc_file.is_file():
            try:
                docs = load_file_documents(str(doc_file))
                if docs:
                    chunks = splitter.split_documents(docs)
                    for c in chunks:
                        if "source" not in c.metadata:
                            c.metadata["source"] = doc_file.name
                    all_chunks.extend(chunks)
            except Exception as e:
                print(f"Skipping {doc_file.name} during index rebuild: {e}")

    embeddings = get_embeddings()

    # Clear old index directory if present
    if os.path.exists(VECTOR_DB_PATH):
        try:
            shutil.rmtree(VECTOR_DB_PATH)
        except Exception as e:
            print(f"Notice: clearing vector db path: {e}")

    if all_chunks:
        vector_db = FAISS.from_documents(all_chunks, embeddings)
        vector_db.save_local(VECTOR_DB_PATH)
        return len(all_chunks)
    return 0


def create_vector_db(file_path: str) -> int:
    """Load an uploaded document and add its chunks to the FAISS vector database."""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found at: {file_path}")

    documents = load_file_documents(file_path)
    if not documents or not any(d.page_content.strip() for d in documents):
        raise ValueError("Could not extract readable text from the document. File may be empty or corrupted.")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=800,
        chunk_overlap=120
    )
    chunks = splitter.split_documents(documents)
    if not chunks:
        raise ValueError("Document was empty or could not be chunked.")

    filename = Path(file_path).name
    for c in chunks:
        if "source" not in c.metadata:
            c.metadata["source"] = filename

    embeddings = get_embeddings()
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    if os.path.exists(VECTOR_DB_PATH):
        try:
            vector_db = FAISS.load_local(
                VECTOR_DB_PATH,
                embeddings,
                allow_dangerous_deserialization=True
            )
            new_db = FAISS.from_documents(chunks, embeddings)
            vector_db.merge_from(new_db)
            vector_db.save_local(VECTOR_DB_PATH)
            return len(chunks)
        except Exception as e:
            print(f"Warning: Could not merge with existing vector DB ({e}), rebuilding.")

    vector_db = FAISS.from_documents(chunks, embeddings)
    vector_db.save_local(VECTOR_DB_PATH)
    return len(chunks)


def search_documents(query: str, k: int = 4) -> List[Dict[str, Any]]:
    """Search uploaded knowledge documents and return formatted results with source metadata."""
    if not os.path.exists(VECTOR_DB_PATH):
        return []

    embeddings = get_embeddings()
    try:
        vector_db = FAISS.load_local(
            VECTOR_DB_PATH,
            embeddings,
            allow_dangerous_deserialization=True
        )
        matched_docs = vector_db.similarity_search(query, k=k)
        results = []
        for d in matched_docs:
            source = d.metadata.get("source", "Document")
            page = d.metadata.get("page")
            meta_str = f"[{source}" + (f", Page {page}" if page else "") + "]"
            results.append({
                "content": d.page_content.strip(),
                "source": source,
                "page": page,
                "citation": meta_str
            })
        return results
    except Exception as e:
        print(f"Error during vector search: {e}")
        return []


def delete_document_file(filename: str) -> bool:
    """Delete a specific document and rebuild the vector index."""
    file_path = DOCS_DIR / filename
    if file_path.exists():
        file_path.unlink()
        rebuild_vector_db_from_scratch()
        return True
    return False


def clear_knowledge_base() -> bool:
    """Delete all documents and vector database."""
    if DOCS_DIR.exists():
        for f in DOCS_DIR.iterdir():
            if f.is_file():
                f.unlink()
    if os.path.exists(VECTOR_DB_PATH):
        shutil.rmtree(VECTOR_DB_PATH, ignore_errors=True)
    return True